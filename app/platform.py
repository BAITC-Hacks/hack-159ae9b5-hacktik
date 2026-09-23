"""Persistent accounts, shared surplus marketplace and retrieval data.

Only demonstration records are seeded. Imported records are retrieved at answer time;
importing a file does not train or change a language model's weights.
"""
from __future__ import annotations

import csv
import hashlib
import hmac
import io
import json
import math
import os
import re
import secrets
import sqlite3
import threading
import time
import uuid
import zipfile
from xml.etree.ElementTree import ParseError
from datetime import datetime, timezone
from pathlib import Path
from typing import Literal

from fastapi import APIRouter, FastAPI, File, Form, HTTPException, Request, Response, UploadFile
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator, model_validator

from .catalog import CartLine, InsufficientStock, PriceChanged, Product, UnknownProduct, demo_products

AUTH_COOKIE = "ekt_auth_session"
SESSION_SECONDS = 7 * 24 * 3600
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def distance_km(lat1: float, lng1: float, lat2: float, lng2: float) -> float:
    a1, a2 = math.radians(lat1), math.radians(lat2)
    delta = math.sin((a2 - a1) / 2) ** 2 + math.cos(a1) * math.cos(a2) * math.sin(math.radians(lng2 - lng1) / 2) ** 2
    return round(6371.0088 * 2 * math.asin(math.sqrt(min(1, max(0, delta)))), 2)


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)

    @field_validator("*", mode="before")
    @classmethod
    def trim_fields(cls, value, info):
        return value.strip() if isinstance(value, str) and info.field_name != "password" else value


class RegisterBody(StrictModel):
    name: str = Field(min_length=2, max_length=100)
    email: str = Field(min_length=5, max_length=254)
    password: str = Field(min_length=8, max_length=128)
    company: str = Field(default="", max_length=160)

    @field_validator("email")
    @classmethod
    def email_valid(cls, value: str) -> str:
        value = value.casefold()
        if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value):
            raise ValueError("Укажите корректный email")
        return value


class LoginBody(StrictModel):
    email: str = Field(min_length=1, max_length=254)
    password: str = Field(min_length=1, max_length=128)


class OfferBody(StrictModel):
    title: str = Field(min_length=3, max_length=180)
    sku: str = Field(default="", max_length=80)
    category: str = Field(min_length=1, max_length=80)
    quantity: float = Field(gt=0, le=1_000_000_000)
    unit: str = Field(default="шт", min_length=1, max_length=20)
    original_price: float = Field(gt=0, le=1_000_000_000)
    price: float = Field(ge=0, le=1_000_000_000)
    lat: float = Field(ge=-90, le=90)
    lng: float = Field(ge=-180, le=180)
    address: str = Field(min_length=3, max_length=300)
    city: str = Field(default="", max_length=100)
    description: str = Field(default="", max_length=3000)
    contact: str = Field(default="", max_length=200)
    expires_at: datetime | None = None

    @model_validator(mode="after")
    def check_discount(self):
        if self.price >= self.original_price:
            raise ValueError("Цена со скидкой должна быть ниже исходной")
        if self.expires_at:
            if self.expires_at.tzinfo is None:
                self.expires_at = self.expires_at.replace(tzinfo=timezone.utc)
            if self.expires_at <= datetime.now(timezone.utc):
                raise ValueError("Дата окончания должна быть в будущем")
        return self


class PersistentCatalog:
    def __init__(self, store: "PlatformStore"):
        self.store = store

    async def search_products(self, query: str) -> list[Product]:
        return [Product.model_validate(item["data"]) for item in self.store.product_records(query)[:5]]

    async def get_product(self, sku: str) -> Product | None:
        with self.store.lock:
            row = self.store.db.execute("SELECT data FROM products WHERE sku=?", (sku.upper(),)).fetchone()
        return Product.model_validate_json(row[0]) if row else None

    async def check_stock(self, sku: str) -> int | None:
        product = await self.get_product(sku)
        return product.stock if product else None

    async def find_alternatives(self, sku: str) -> list[Product]:
        original = await self.get_product(sku)
        if not original:
            return []
        products = [Product.model_validate(p["data"]) for p in self.store.product_records()]
        # Missing compatibility data never establishes technical equivalence.
        return [p for p in products if original.compatibility and p.sku != original.sku
                and p.stock > 0 and p.category == original.category
                and p.compatibility == original.compatibility][:3]

    async def add_checked(self, session_id: str, sku: str, quantity: int, locale: str,
                          expected_price_kzt: int) -> CartLine:
        if quantity <= 0:
            raise ValueError("Quantity must be positive")
        with self.store.lock, self.store.db:
            # BEGIN IMMEDIATE serializes stock reservation across application processes.
            self.store.db.execute("BEGIN IMMEDIATE")
            row = self.store.db.execute("SELECT data FROM products WHERE sku=?", (sku.upper(),)).fetchone()
            if not row:
                raise UnknownProduct()
            product = Product.model_validate_json(row[0])
            if product.price_kzt != expected_price_kzt:
                raise PriceChanged(product.price_kzt)
            if product.stock < quantity:
                raise InsufficientStock(product.stock)
            product.stock -= quantity
            line = CartLine(sku=product.sku, name=product.name(locale), quantity=quantity,
                            unit_price_kzt=product.price_kzt)
            self.store.db.execute("UPDATE products SET data=? WHERE sku=?", (product.model_dump_json(), product.sku))
            self.store.db.execute("INSERT INTO carts(session_id,data) VALUES(?,?)", (session_id, line.model_dump_json()))
            return line

    async def get_cart(self, session_id: str) -> list[CartLine]:
        with self.store.lock:
            rows = self.store.db.execute("SELECT data FROM carts WHERE session_id=? ORDER BY id", (session_id,)).fetchall()
        return [CartLine.model_validate_json(row[0]) for row in rows]

    async def checkout_url(self, session_id: str) -> str:
        return "/demo-cart"


class PlatformStore:
    def __init__(self, db_path: str | Path | None = None):
        path = str(db_path or os.getenv("EKT_DATABASE_PATH", str(Path(__file__).resolve().parent.parent / "data" / "ekt.sqlite3")))
        if path != ":memory:":
            Path(path).parent.mkdir(parents=True, exist_ok=True)
        self.lock = threading.RLock()
        self.db = sqlite3.connect(path, check_same_thread=False, timeout=10)
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.execute("PRAGMA foreign_keys=ON")
        with self.db:
            self.db.executescript("""
                CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY,name TEXT NOT NULL,email TEXT UNIQUE NOT NULL,
                    password_hash TEXT NOT NULL,company TEXT NOT NULL,role TEXT NOT NULL,created_at TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS auth_sessions(token_hash TEXT PRIMARY KEY,user_id TEXT NOT NULL REFERENCES users(id),expires_at REAL NOT NULL);
                CREATE TABLE IF NOT EXISTS products(sku TEXT PRIMARY KEY,data TEXT NOT NULL,is_demo INTEGER NOT NULL DEFAULT 0,source TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS offers(id TEXT PRIMARY KEY,owner_id TEXT NOT NULL,data TEXT NOT NULL,is_demo INTEGER NOT NULL DEFAULT 0,created_at TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS knowledge(id TEXT PRIMARY KEY,title TEXT NOT NULL,content TEXT NOT NULL,source TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS imports(id TEXT PRIMARY KEY,name TEXT NOT NULL,kind TEXT NOT NULL,rows INTEGER NOT NULL,imported_at TEXT NOT NULL,user_id TEXT NOT NULL);
                CREATE TABLE IF NOT EXISTS carts(id INTEGER PRIMARY KEY AUTOINCREMENT,session_id TEXT NOT NULL,data TEXT NOT NULL);
                CREATE INDEX IF NOT EXISTS carts_session ON carts(session_id);
                CREATE TABLE IF NOT EXISTS metadata(key TEXT PRIMARY KEY,value TEXT NOT NULL);
            """)
            self.db.execute("BEGIN IMMEDIATE")
            seed_demo = os.getenv("EKT_SEED_DEMO", "true").casefold() not in {"0", "false", "no"}
            if seed_demo and not self.db.execute("SELECT 1 FROM metadata WHERE key='seeded'").fetchone():
                for product in demo_products():
                    self.db.execute("INSERT OR IGNORE INTO products VALUES(?,?,1,'Демонстрационный каталог')", (product.sku, product.model_dump_json()))
                for index, fields in enumerate([
                    ("Кабель ВВГнг 3×2.5", "cable", 240, "м", 580, 410, 43.2389, 76.8897, "Алматы, район Толе би / Розыбакиева"),
                    ("Автоматы C16, 1P", "breaker", 60, "шт", 3200, 2240, 43.2567, 76.9286, "Алматы, район проспекта Райымбека"),
                    ("LED лампы E27, 12 Вт", "lamp", 120, "шт", 1150, 690, 43.2207, 76.9292, "Алматы, район Бостандык"),
                ]):
                    title, category, quantity, unit, original, price, lat, lng, address = fields
                    offer = OfferBody(title=title, category=category, quantity=quantity, unit=unit, original_price=original,
                                      price=price, lat=lat, lng=lng, address=address, city="Алматы",
                                      description="Демонстрационное предложение. Товары и адрес условные; продажа не осуществляется.")
                    self.db.execute("INSERT INTO offers VALUES(?,?,?,1,?)", (f"demo-{index + 1}", "demo", offer.model_dump_json(), now_iso()))
                self.db.execute("INSERT INTO metadata VALUES('seeded','1')")
        self.catalog = PersistentCatalog(self)

    def close(self):
        with self.lock:
            self.db.close()

    @staticmethod
    def public_user(row) -> dict:
        return {key: row[key] for key in ("id", "name", "email", "role")}

    def user_for_token(self, token: str | None):
        if not token or len(token) > 200:
            return None
        with self.lock:
            row = self.db.execute("SELECT u.* FROM users u JOIN auth_sessions s ON u.id=s.user_id WHERE s.token_hash=? AND s.expires_at>?",
                                  (hashlib.sha256(token.encode()).hexdigest(), time.time())).fetchone()
        return self.public_user(row) if row else None

    def new_session(self, user_id: str) -> str:
        token = secrets.token_urlsafe(32)
        with self.lock, self.db:
            self.db.execute("DELETE FROM auth_sessions WHERE expires_at<?", (time.time(),))
            self.db.execute("INSERT INTO auth_sessions VALUES(?,?,?)", (hashlib.sha256(token.encode()).hexdigest(), user_id, time.time() + SESSION_SECONDS))
        return token

    def revoke(self, token: str | None):
        if token:
            with self.lock, self.db:
                self.db.execute("DELETE FROM auth_sessions WHERE token_hash=?", (hashlib.sha256(token.encode()).hexdigest(),))

    def product_records(self, query: str = "") -> list[dict]:
        with self.lock:
            rows = self.db.execute("SELECT * FROM products ORDER BY sku").fetchall()
        records = [{"data": json.loads(row["data"]), "is_demo": bool(row["is_demo"]), "source": row["source"]} for row in rows]
        return self._rank(records, query, lambda r: json.dumps(r["data"], ensure_ascii=False))

    @staticmethod
    def _rank(items: list, query: str, text_for) -> list:
        words = set(re.findall(r"[\w]+", query.casefold()))
        if not words:
            return items
        scores = [(sum(word in text_for(item).casefold() for word in words), index, item) for index, item in enumerate(items)]
        return [item for score, _, item in sorted(scores, key=lambda x: (-x[0], x[1])) if score]

    def offers(self, query: str = "", category: str = "", lat: float | None = None, lng: float | None = None,
               radius_km: float | None = None, owner_id: str | None = None) -> list[dict]:
        with self.lock:
            rows = self.db.execute("SELECT o.*,u.name AS owner_name FROM offers o LEFT JOIN users u ON u.id=o.owner_id ORDER BY o.created_at DESC").fetchall()
        items = []
        for row in rows:
            item = json.loads(row["data"])
            if owner_id and row["owner_id"] != owner_id:
                continue
            if category and item["category"] != category:
                continue
            if item.get("expires_at") and datetime.fromisoformat(item["expires_at"]) <= datetime.now(timezone.utc):
                continue
            item.update(id=row["id"], owner_id=row["owner_id"], owner_name=row["owner_name"] or "Демо-склад",
                        is_demo=bool(row["is_demo"]), created_at=row["created_at"],
                        discount_percent=round((1 - item["price"] / item["original_price"]) * 100))
            if lat is not None and lng is not None:
                item["distance_km"] = distance_km(lat, lng, item["lat"], item["lng"])
                if radius_km is not None and item["distance_km"] > radius_km:
                    continue
            items.append(item)
        items = self._rank(items, query, lambda x: " ".join(str(x.get(key, "")) for key in ("title", "sku", "category", "city", "address", "description")))
        if lat is not None and lng is not None:
            items.sort(key=lambda item: item["distance_km"])
        return items

    def status(self) -> dict:
        with self.lock:
            product_count = self.db.execute("SELECT COUNT(*),COALESCE(SUM(is_demo),0) FROM products").fetchone()
            knowledge = self.db.execute("SELECT COUNT(*) FROM knowledge").fetchone()[0]
            sources = [dict(row) for row in self.db.execute("SELECT name,kind,rows,imported_at FROM imports ORDER BY imported_at DESC LIMIT 30")]
        offers = self.offers()
        return {"products": product_count[0], "demo_products": product_count[1], "real_products": product_count[0] - product_count[1],
                "knowledge": knowledge, "offers": len(offers), "real_offers": sum(not o["is_demo"] for o in offers),
                "sources": sources, "model_trained": False, "mode": "retrieval"}

    def context(self, query: str, lat: float | None = None, lng: float | None = None) -> dict:
        with self.lock:
            rows = [dict(row) for row in self.db.execute("SELECT title,content,source FROM knowledge")]
        knowledge = self._rank(rows, query, lambda row: row["title"] + " " + row["content"])[:6]
        for row in knowledge:
            row["content"] = row["content"][:2400]
        products = [dict(row["data"], is_demo=row["is_demo"], source=row["source"]) for row in self.product_records(query)[:8]]
        offers = self.offers(query, lat=lat, lng=lng)[:8]
        # General location questions often have no product tokens; nearby public offers remain useful.
        if not offers and (lat is not None or re.search(r"скид|избыт|остат|рядом|близ|карт|discount|nearby", query, re.I)):
            offers = self.offers(lat=lat, lng=lng)[:8]
        return {"products": products, "offers": offers, "knowledge": knowledge, "status": self.status()}

    def import_records(self, content: bytes, filename: str, kind: str, user_id: str) -> dict:
        name = Path(filename.replace("\\", "/")).name[:150]
        if kind not in {"products", "knowledge"}:
            raise ValueError("Тип импорта должен быть products или knowledge")
        if not content or len(content) > MAX_FILE_BYTES:
            raise ValueError("Файл должен быть непустым и не больше 5 МБ")
        suffix = Path(name).suffix.casefold()
        if suffix == ".json":
            payload = json.loads(content.decode("utf-8-sig"))
            records = payload.get(kind, payload.get("items")) if isinstance(payload, dict) else payload
        elif suffix == ".csv":
            raw = content.decode("utf-8-sig")
            try:
                dialect = csv.Sniffer().sniff(raw[:8192], delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            records = list(csv.DictReader(io.StringIO(raw), dialect=dialect))
        elif suffix == ".xlsx":
            # Reject ZIP bombs before the spreadsheet parser decompresses a workbook.
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(entry.file_size for entry in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError("Распакованный XLSX превышает 40 МБ")
            from openpyxl import load_workbook
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
            try:
                if not book.active or (book.active.max_column or 0) > 100:
                    raise ValueError("XLSX должен содержать лист с не более чем 100 столбцами")
                rows = book.active.iter_rows(values_only=True)
                headers = [str(value).strip() if value is not None else "" for value in next(rows, [])]
                records = []
                for row_number, row in enumerate(rows, 1):
                    if row_number > MAX_ROWS:
                        raise ValueError("Не больше 10 000 строк за один импорт")
                    if any(value is not None for value in row):
                        records.append(dict(zip(headers, row)))
            finally:
                book.close()
        else:
            raise ValueError("Поддерживаются CSV (UTF-8), JSON и XLSX. Экспортируйте базу данных в один из этих форматов")
        if not isinstance(records, list) or not records or len(records) > MAX_ROWS:
            raise ValueError("Ожидается от 1 до 10 000 записей")
        parsed, seen = [], set()
        for number, row in enumerate(records, 1):
            if not isinstance(row, dict):
                raise ValueError(f"Строка {number}: ожидается объект с полями")
            try:
                if kind == "products":
                    row = {str(k).strip(): v for k, v in row.items() if k is not None}
                    sku = str(row.get("sku") or "").strip().upper()
                    name_ru = str(row.get("name_ru") or row.get("name") or "").strip()
                    if not sku or len(sku) > 80 or not name_ru or len(name_ru) > 300:
                        raise ValueError("нужны sku (до 80 символов) и name_ru (до 300)")
                    if sku in seen:
                        raise ValueError("артикул повторяется в файле")
                    seen.add(sku)
                    data = {key: row[key] for key in ("category", "price_kzt", "stock", "certificate_url") if row.get(key) not in (None, "")}
                    if "category" not in data or len(str(data["category"])) > 80:
                        raise ValueError("нужна category (до 80 символов)")
                    data.update(sku=sku, name_ru=name_ru, name_kk=str(row.get("name_kk") or name_ru)[:300])
                    for key in ("specifications", "compatibility"):
                        value = row.get(key) or {}
                        if isinstance(value, str):
                            value = json.loads(value)
                        if not isinstance(value, dict) or len(json.dumps(value, ensure_ascii=False)) > 10000:
                            raise ValueError(f"{key} должен быть небольшим JSON-объектом")
                        data[key] = {str(k): str(v) for k, v in value.items()}
                    parsed.append(Product.model_validate(data))
                else:
                    title = str(row.get("title") or row.get("question") or "").strip()
                    text = str(row.get("content") or row.get("answer") or row.get("text") or "").strip()
                    if not title or not text or len(title) > 300 or len(text) > 20000:
                        raise ValueError("нужны title/question (до 300 символов) и content/answer (до 20 000)")
                    parsed.append((title, text))
            except (ValidationError, ValueError, TypeError) as exc:
                raise ValueError(f"Строка {number}: {str(exc)[:250]}") from exc
        # Validation is complete before the transaction; a bad row never partially imports a file.
        with self.lock, self.db:
            for record in parsed:
                if kind == "products":
                    self.db.execute("INSERT INTO products VALUES(?,?,0,?) ON CONFLICT(sku) DO UPDATE SET data=excluded.data,is_demo=0,source=excluded.source",
                                    (record.sku, record.model_dump_json(), name))
                else:
                    key = hashlib.sha256((name + "\0" + record[0]).encode()).hexdigest()
                    self.db.execute("INSERT INTO knowledge VALUES(?,?,?,?) ON CONFLICT(id) DO UPDATE SET content=excluded.content,source=excluded.source",
                                    (key, record[0], record[1], name))
            imported_at = now_iso()
            self.db.execute("INSERT INTO imports VALUES(?,?,?,?,?,?)", (uuid.uuid4().hex, name, kind, len(parsed), imported_at, user_id))
        return {"name": name, "kind": kind, "rows": len(parsed), "imported_at": imported_at, "status": self.status()}


def password_hash(password: str, salt: str | None = None) -> str:
    salt = salt or secrets.token_hex(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), bytes.fromhex(salt), 310000).hex()
    return f"pbkdf2_sha256$310000${salt}${digest}"


def mount_platform(app: FastAPI, db_path: str | Path | None = None) -> PlatformStore:
    store = PlatformStore(db_path)
    app.state.platform = store
    router = APIRouter(prefix="/api")
    login_attempts: dict[str, list[float]] = {}

    def require_user(request: Request):
        user = store.user_for_token(request.cookies.get(AUTH_COOKIE))
        if not user:
            raise HTTPException(401, "Войдите в аккаунт менеджера")
        return user

    def throttle(request: Request):
        key = request.client.host if request.client else "unknown"
        stamp = time.time()
        with store.lock:
            values = [value for value in login_attempts.get(key, []) if value > stamp - 600]
            if len(values) >= 30:
                raise HTTPException(429, "Слишком много попыток. Повторите через 10 минут")
            login_attempts[key] = values + [stamp]
            if len(login_attempts) > 10000:
                for old in list(login_attempts):
                    if not login_attempts[old] or login_attempts[old][-1] < stamp - 600:
                        del login_attempts[old]

    def set_session(request: Request, response: Response, user_id: str):
        store.revoke(request.cookies.get(AUTH_COOKIE))
        response.delete_cookie("ekt_assistant_session", path="/")
        response.set_cookie(AUTH_COOKIE, store.new_session(user_id), httponly=True, samesite="strict",
                            secure=request.url.scheme == "https", max_age=SESSION_SECONDS, path="/")

    @router.get("/auth/me")
    def me(request: Request):
        return {"user": store.user_for_token(request.cookies.get(AUTH_COOKIE))}

    @router.post("/auth/register", status_code=201)
    def register(body: RegisterBody, request: Request, response: Response):
        throttle(request)
        user = {"id": uuid.uuid4().hex, "name": body.name, "email": body.email, "role": "manager"}
        encoded_password = password_hash(body.password)
        try:
            with store.lock, store.db:
                store.db.execute("INSERT INTO users VALUES(?,?,?,?,?,?,?)", (user["id"], body.name, body.email,
                                 encoded_password, body.company, "manager", now_iso()))
        except sqlite3.IntegrityError as exc:
            raise HTTPException(409, "Такой email уже зарегистрирован") from exc
        set_session(request, response, user["id"])
        return {"user": user}

    @router.post("/auth/login")
    def login(body: LoginBody, request: Request, response: Response):
        throttle(request)
        with store.lock:
            row = store.db.execute("SELECT * FROM users WHERE email=?", (body.email.casefold(),)).fetchone()
        # Always perform a password derivation, including for unknown email addresses.
        expected = row["password_hash"] if row else ""
        salt = expected.split("$")[2] if expected else "0" * 32
        actual = password_hash(body.password, salt)
        if not row or not hmac.compare_digest(actual, expected):
            raise HTTPException(401, "Неверный email или пароль")
        set_session(request, response, row["id"])
        return {"user": store.public_user(row)}

    @router.post("/auth/logout")
    def logout(request: Request, response: Response):
        store.revoke(request.cookies.get(AUTH_COOKIE))
        response.delete_cookie(AUTH_COOKIE, path="/")
        response.delete_cookie("ekt_assistant_session", path="/")
        return {"ok": True}

    @router.get("/offers")
    def list_offers(request: Request, q: str = "", category: str = "", lat: float | None = None,
                    lng: float | None = None, radius_km: float | None = None, mine: bool = False):
        if (lat is None) != (lng is None):
            raise HTTPException(422, "Укажите обе координаты lat и lng")
        if lat is not None and (not math.isfinite(lat) or not math.isfinite(lng) or not -90 <= lat <= 90 or not -180 <= lng <= 180):
            raise HTTPException(422, "Некорректные координаты")
        if radius_km is not None and (lat is None or not math.isfinite(radius_km) or not 0 < radius_km <= 20050):
            raise HTTPException(422, "Радиус должен быть от 0 до 20 050 км; укажите координаты")
        user = require_user(request) if mine else None
        items = store.offers(q[:300], category[:80], lat, lng, radius_km, user["id"] if user else None)
        return {"items": items, "total": len(items)}

    @router.post("/offers", status_code=201)
    def create_offer(body: OfferBody, request: Request):
        user = require_user(request)
        offer_id = uuid.uuid4().hex
        with store.lock, store.db:
            store.db.execute("INSERT INTO offers VALUES(?,?,?,0,?)", (offer_id, user["id"], body.model_dump_json(), now_iso()))
        return {"item": next(item for item in store.offers(owner_id=user["id"]) if item["id"] == offer_id)}

    @router.put("/offers/{offer_id}")
    def update_offer(offer_id: str, body: OfferBody, request: Request):
        user = require_user(request)
        with store.lock, store.db:
            changed = store.db.execute("UPDATE offers SET data=? WHERE id=? AND owner_id=? AND is_demo=0", (body.model_dump_json(), offer_id, user["id"])).rowcount
        if not changed:
            raise HTTPException(404, "Ваше предложение не найдено")
        return {"item": next(item for item in store.offers(owner_id=user["id"]) if item["id"] == offer_id)}

    @router.delete("/offers/{offer_id}")
    def delete_offer(offer_id: str, request: Request):
        user = require_user(request)
        with store.lock, store.db:
            changed = store.db.execute("DELETE FROM offers WHERE id=? AND owner_id=? AND is_demo=0", (offer_id, user["id"])).rowcount
        if not changed:
            raise HTTPException(404, "Ваше предложение не найдено")
        return {"ok": True}

    @router.get("/products")
    def products(q: str = ""):
        items = [dict(row["data"], is_demo=row["is_demo"], source=row["source"]) for row in store.product_records(q[:300])[:200]]
        return {"items": items, "total": len(items)}

    @router.get("/data/status")
    def status():
        return store.status()

    @router.post("/data/import")
    async def import_data(request: Request, file: UploadFile = File(...), kind: Literal["products", "knowledge"] = Form(...)):
        user = require_user(request)
        content = await file.read(MAX_FILE_BYTES + 1)
        await file.close()
        try:
            # Spreadsheet parsing is blocking; keep it off the server's event loop.
            from starlette.concurrency import run_in_threadpool
            return await run_in_threadpool(store.import_records, content, file.filename or "upload", kind, user["id"])
        except (ValueError, UnicodeError, csv.Error, zipfile.BadZipFile, KeyError, TypeError, ParseError, RecursionError) as exc:
            raise HTTPException(422, str(exc)[:500]) from exc

    app.include_router(router)
    return store
